"""Copy of the blank sheet with everything except the written answers filled in.

python fill_data_docx.py <blank HW2_Answers.docx> <results dir> <out HW2_Answers_data.docx> [HW2 folder]

With the HW2 folder: also screenshots, blurred images, header, P1/P2, extra-credit AI table and NCU text.
"""
import statistics
import sys
from collections import defaultdict
from pathlib import Path

import docx
from docx.shared import Inches, Pt

KMAP = {"global": "global", "static": "static SMEM", "dynamic": "dynamic SMEM"}


def timing(path):
    """median of reps 2..N per (img, kernel); cpu t3 per img"""
    reps, cpu = defaultdict(list), {}
    for line in Path(path).read_text(encoding="utf-8-sig").splitlines():
        if not line.startswith("RESULT"):
            continue
        kv = dict(x.split("=", 1) for x in line.split()[1:])
        if kv["kernel"] == "cpu":
            cpu[int(kv["img"])] = float(kv["t3_ms"])
        elif int(kv["rep"]) > 1:
            reps[(int(kv["img"]), kv["kernel"])].append((float(kv["t1_ms"]), float(kv["t2_ms"])))
    med = {k: (statistics.median(a for a, _ in v), statistics.median(b for _, b in v)) for k, v in reps.items()}
    return med, cpu


UNIT = {"byte": 1, "Kbyte": 1e3, "Mbyte": 1e6, "Gbyte": 1e9, "cycle": 1, "inst": 1, "sector": 1,
        "Ghz": 1e9, "Mhz": 1e6, "Khz": 1e3, "hz": 1, "cycle/nsecond": 1e9, "cycle/usecond": 1e6,
        "cycle/msecond": 1e3, "cycle/second": 1, "ns": 1e-9, "nsecond": 1e-9, "us": 1e-6, "usecond": 1e-6,
        "ms": 1e-3, "msecond": 1e-3}


def ncu(path):
    """metrics from the plain ncu text output (same run as the pasted CLI output), in base units"""
    m = {}
    for line in Path(path).read_text(encoding="utf-8-sig").splitlines():
        parts = line.split()
        if len(parts) == 3 and "__" in parts[0] and parts[1] in UNIT:
            m[parts[0]] = float(parts[2].replace(",", "")) * UNIT[parts[1]]
    return m


def put(cell, text):
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    r = p.add_run(text)
    r.font.size = Pt(8.5)


def find(d, first_header, second_header):
    for t in d.tables:
        h = [c.text for c in t.rows[0].cells]
        if h[0] == first_header and h[1].startswith(second_header):
            return t
    raise LookupError(first_header)


def main(blank, res, out, shots=None):
    res = Path(res)
    rel, cpu_rel = timing(res / "timing_release.txt")
    dbg, cpu_dbg = timing(res / "timing_debug.txt")
    d = docx.Document(blank)

    t = find(d, "Image", "Kernel")
    for row in t.rows[1:]:
        img = int(row.cells[0].text.split()[0][1:])
        k = next(k for k, v in KMAP.items() if v == row.cells[1].text)
        for i, v in enumerate((*dbg[(img, k)], *rel[(img, k)]), start=2):
            put(row.cells[i], f"{v:.3f}")

    t = find(d, "Kernel", "Image")
    for row in t.rows[1:]:
        k = "global" if row.cells[0].text == "global" else "shared"
        img = int(row.cells[1].text[1:])
        m = ncu(res / "ncu" / f"ncu_{k}_img{img}.txt")
        cyc, ps = m["smsp__cycles_elapsed.sum"], m["smsp__cycles_elapsed.sum.per_second"]
        inst, db = m["smsp__inst_executed.sum"], m["dram__bytes.sum"]
        vals = [f"{cyc:,.0f}", f"{ps:.4g}", f"{inst:,.0f}", f"{db:,.0f}", f"{cyc / ps:.4e}", f"{inst / db:.2f}"]
        for i, v in enumerate(vals, start=2):
            put(row.cells[i], v)

    t = find(d, "Image", "CPU T3")
    for row in t.rows[1:]:
        img = int(row.cells[0].text.split()[0][1:])
        g = [rel[(img, k)][1] for k in ("global", "static", "dynamic")]
        vals = [f"{cpu_dbg[img]:.2f}", f"{cpu_rel[img]:.2f}"] + [f"{v:.3f}" for v in g]
        for i, v in enumerate(vals, start=1):
            put(row.cells[i], v)

    if shots:
        fill_debugger(d, Path(shots))
        fill_static(d, Path(shots), res.parent / "fig")
    d.save(out)


def box_after(d, label_prefix):
    """the 1x1 answer box that follows the bold label paragraph"""
    body = list(d.element.body)
    for i, el in enumerate(body):
        if el.tag.endswith("}p") and "".join(el.itertext()).startswith(label_prefix):
            docx.text.paragraph.Paragraph(el, d._body).paragraph_format.keep_with_next = True
            nxt = body[i + 1]
            return docx.table.Table(nxt, d._body).rows[0].cells[0]
    raise LookupError(label_prefix)


def put_pictures(cell, paths, width_in):
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    for i, path in enumerate(paths):
        if i:
            p.add_run("  ")
        p.add_run().add_picture(str(path), width=Inches(width_in))


def fill_debugger(d, hw2_dir):
    """Nsight screenshots and blurred images #1-#3"""
    put_pictures(box_after(d, "Screen capture 1"), [hw2_dir / "screenshot_1.png"], 6.2)
    put_pictures(box_after(d, "Screen capture 2"), [hw2_dir / "screenshot_2.png"], 6.2)
    imgs = [hw2_dir / "output" / f"hw2_outimage{i}_global.png" for i in (1, 2, 3)]
    put_pictures(box_after(d, "Blurred output images"), imgs, 1.95)


def put_lines(cell, lines, size=8.5, mono=False):
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    for i, line in enumerate(lines):
        r = p.add_run(line)
        r.font.size = Pt(size)
        if mono:
            r.font.name = "Consolas"
        if i < len(lines) - 1:
            r.add_break()


def fill_row(t, label, vals, col0=1):
    for row in t.rows[1:]:
        if row.cells[0].text == label:
            for i, v in enumerate(vals, start=col0):
                put(row.cells[i], v)
            return row
    raise LookupError(label)


def tables_after(d, label_prefix, header0):
    """tables with header cell 0 == header0 that follow each paragraph starting with label_prefix"""
    out, armed = [], False
    for el in d.element.body:
        if el.tag.endswith("}p") and "".join(el.itertext()).startswith(label_prefix):
            armed = True
        elif armed and el.tag.endswith("}tbl"):
            t = docx.table.Table(el, d._body)
            if t.rows[0].cells[0].text == header0:
                out.append(t)
                armed = False
    return out


def ncu_metrics(txt):
    lines = Path(txt).read_text(encoding="utf-8-sig").splitlines()
    head = next(l.strip() for l in lines if "gaussBlurKernel" in l and "(" in l and "CC" in l)
    start = next(i for i, l in enumerate(lines) if "Section: Command line profiler metrics" in l)
    body = [l.rstrip()[4:] for l in lines[start + 1:] if l.strip()]
    body = [l for l in body if not l.strip().startswith("---")]
    return [head] + body


def fill_static(d, hw2_dir, fig_dir):
    """header, P1/P2 tables and plots, extra-credit AI table, NCU text"""
    t = d.tables[0]
    hw1 = docx.Document(str(hw2_dir.parent / "HW1" / "HW1_Answers.docx")).tables[0]
    info = {r.cells[0].text: r.cells[1].text for r in hw1.rows[1:]}
    # tell the grader the numbers come from this PC (RTX 3050), not the AWS T4
    info["Build environment"] = info["Build environment"].replace(
        "code generation compute_75,sm_75 (PTX JIT on device)", "code generation compute_75,sm_75;compute_86,sm_86")
    info["Build environment"] = (
        "Local PC, not the AWS WorkSpaces Tesla T4: NVIDIA GeForce RTX 3050 8 GB (compute capability 8.6), driver 617.42; "
        "Visual Studio Community 2026 18.10.3 (MSVC v145), CUDA 13.4 (nvcc V13.4.92), x64. "
        "Timing and NCU runs: nvcc command-line builds, Release -O3 and Debug -G, code generation compute_75,sm_75 "
        "(PTX JIT-compiled by the driver on the RTX 3050). Debugger screenshots: Visual Studio Debug build, "
        "compute_75,sm_75;compute_86,sm_86.")
    for k in ("Name", "UW NetID", "Build environment"):
        fill_row(t, k, [info[k]])

    # P1
    boxes = [docx.table.Table(el, d._body).rows[0].cells[0] for el in _boxes_after(d, "Algorithm equation and term definitions")]
    p1_eq = [
        ["B[j][i] = A[i][j] for 0 <= i, j < N (FP32, out-of-place).",
         "Read: every element of A once (4 B). Write: every element of B once (4 B). No arithmetic."],
        ["Input x of length N (one row; R rows multiply every term by R, AI unchanged).",
         "Pass 1: m = max_i x_i   (N-1 compares, read x)",
         "Pass 2: e_i = exp(x_i - m), s = sum_i e_i   (N subtracts + N exp x 4 FLOP + N-1 adds; read x, write e)",
         "Pass 3: y_i = e_i / s   (N divides, counted as 1 FLOP each; read e, write y)",
         "Unfused: each pass reads its input from DRAM and the intermediate e is stored to DRAM."],
        ["Out[n,co,h,w] = sum_{ci,kh,kw} In[n,ci,h+kh,w+kw] * Wt[co,ci,kh,kw]",
         "Direct conv, stride 1, same padding, C_in = C_out = C, batch N, image W x H, kernel K x K.",
         "Each output: C*K^2 multiply-accumulates = C*K^2 FMA = 2*C*K^2 FLOP.",
         "Each input, weight and output element touches DRAM once (registers/cache free)."],
    ]
    for cell, lines in zip(boxes, p1_eq):
        put_lines(cell, lines)
    p1_rows = [
        [["read A", "0", "4N²", "—"], ["write B = Aᵀ", "0", "4N²", "—"], ["Total", "0", "8N²", "0"]],
        [["pass 1: max", "N − 1", "4N (read x)", "—"], ["pass 2: exp, sum", "N + 4N + (N − 1)", "4N read x + 4N write e", "—"],
         ["pass 3: divide", "N", "4N read e + 4N write y", "—"], ["Total", "≈ 8N", "20N", "0.4"]],
        [["multiply-accumulate (FMA)", "2·N·C²·H·W·K²", "—", "—"], ["read input + weights", "—", "4·N·C·H·W + 4·C²·K²", "—"],
         ["write output", "—", "4·N·C·H·W", "—"],
         ["Total", "2·N·C²·H·W·K²", "8·N·C·H·W + 4·C²·K²", "2NC²HWK² / (8NCHW + 4C²K²) ≈ C·K²/4"]],
    ]
    for t, rows in zip(tables_after(d, "Algorithm equation and term definitions", "algorithm"), p1_rows):
        for row, vals in zip(t.rows[1:], rows):
            for i, v in enumerate(vals):
                put(row.cells[i], v)
    order = ["FLOPs: none; Mem bytes: O(N²); AI = 0",
             "FLOPs: O(N); Mem bytes: O(N) per row; AI: O(1) = 0.4",
             "FLOPs: O(N·C²·H·W·K²); Mem bytes: O(N·C·H·W + C²·K²); AI: O(C·K²)"]
    for el, text in zip([e for e in d.element.body if e.tag.endswith("}p")
                         and "".join(e.itertext()).startswith("Order of complexity:")], order):
        docx.text.paragraph.Paragraph(el, d._body).runs[1].text = " " + text

    # P2
    t = find(d, "GPU", "Peak FP32")
    fill_row(t, "Quadro RTX 6000 (TU102)", ["16.3", "672", "24.3"])
    fill_row(t, "H100 SXM5", ["67", "3350", "20.0"])
    put_lines(box_after(d, "Chosen values for the variable terms"), [
        "Transpose: N = 4096 → FLOPs 0, Mem 134.2 MB, AI 0 (pinned to the left edge of the log axis).",
        "SoftMax: N = 4096 per row, 1024 rows → 33.6 MFLOP, 83.9 MB, AI 0.4.",
        "Convolution: N = 32, C = 64, H = W = 224, K = 3 → 118.4 GFLOP, 822 MB, AI 144.",
        "Regimes: transpose and softmax are memory-bound and conv is compute-bound on both GPUs; none is in the "
        "transitional regime near the ridge (24.3 / 20.0 FLOP/B). Transpose and softmax have a fixed AI, while conv's "
        "AI ≈ C·K²/4 moves with the chosen values (C = 3, K = 3 gives 6.75, memory-bound)."])
    for box, png in zip(_boxes_after(d, "Roofline plot"), ["roofline_quadro.png", "roofline_h100.png"]):
        put_pictures(docx.table.Table(box, d._body).rows[0].cells[0], [fig_dir / png], 6.2)
    # extra credit AI table
    t = find(d, "Kernel", "FLOPs per output pixel")
    fill_row(t, "Global memory", ["≈ 163 (81 FMUL + 81 FADD + 1 FDIV)", "406 (81 pixels + 81×4 coeffs + 1 write), no cache", "≈ 0.40"])
    fill_row(t, "Static shared memory", ["≈ 163", "≈ 83.3 (81 + 324/256 + 1), no cache", "≈ 1.96"])

    # NCU text
    ncu_dir = Path(__file__).resolve().parent.parent / "results" / "ncu"
    lines = []
    for k in ("global", "shared"):
        for i in range(1, 5):
            lines += [f"== ncu ./hw2_rel_{k}_img{i}.exe"] + ncu_metrics(ncu_dir / f"ncu_{k}_img{i}.txt") + [""]
    cell = box_after(d, "NCU CLI output")
    put_lines(cell, lines[:-1], size=6.5, mono=True)
    # longer than a page: let this one box break across pages instead of leaving the label alone on a page
    for cs in cell._tc.getparent().iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit"):
        cs.getparent().remove(cs)

    # the spacer paragraph after the last table can spill onto a blank final page
    body = d.element.body
    while body[-2].tag.endswith("}p") and not "".join(body[-2].itertext()).strip():
        body.remove(body[-2])
    # Word needs a paragraph after a final table: make it 1 pt with no spacing
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    body.insert(len(body) - 1, docx.oxml.parse_xml(
        f'<w:p xmlns:w="{W}"><w:pPr><w:spacing w:before="0" w:after="0" w:line="20" w:lineRule="exact"/>'
        f'<w:rPr><w:sz w:val="2"/></w:rPr></w:pPr></w:p>'))


def _boxes_after(d, label_prefix):
    body = list(d.element.body)
    return [body[i + 1] for i, el in enumerate(body)
            if el.tag.endswith("}p") and "".join(el.itertext()).startswith(label_prefix)]


if __name__ == "__main__":
    main(*sys.argv[1:5])
