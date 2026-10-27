"""Blank HW2 answer sheet in the HW1_Answers.docx style.

python make_answers_docx.py <HW1_Answers.docx (style base)> <out HW2_Answers.docx>
"""
import copy
import sys
from xml.sax.saxutils import escape

import docx
from docx.oxml import parse_xml
from docx.shared import Pt

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
TBL_PR = """<w:tblPr><w:tblW w:w="9360" w:type="dxa"/><w:tblBorders>
<w:top w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:left w:val="single" w:sz="4" w:space="0" w:color="auto"/>
<w:bottom w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:right w:val="single" w:sz="4" w:space="0" w:color="auto"/>
<w:insideH w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="auto"/>
</w:tblBorders><w:tblCellMar><w:left w:w="10" w:type="dxa"/><w:right w:w="10" w:type="dxa"/></w:tblCellMar>
<w:tblLook w:val="04A0" w:firstRow="1" w:lastRow="0" w:firstColumn="1" w:lastColumn="0" w:noHBand="0" w:noVBand="1"/></w:tblPr>"""
BORDERS = """<w:tcBorders><w:top w:val="single" w:sz="6" w:space="0" w:color="808080"/><w:left w:val="single" w:sz="6" w:space="0" w:color="808080"/>
<w:bottom w:val="single" w:sz="6" w:space="0" w:color="808080"/><w:right w:val="single" w:sz="6" w:space="0" w:color="808080"/></w:tcBorders>"""


def _cell(width, text, header=False, size=None):
    shd = '<w:shd w:val="clear" w:color="auto" w:fill="E8E8E8"/>' if header else ""
    rpr = "<w:rPr>" + ("<w:b/><w:bCs/>" if header else "") + (f'<w:sz w:val="{size}"/><w:szCs w:val="{size}"/>' if size else "") + "</w:rPr>"
    run = f'<w:r>{rpr}<w:t xml:space="preserve">{escape(text)}</w:t></w:r>' if text else ""
    return (f'<w:tc><w:tcPr><w:tcW w:w="{width}" w:type="dxa"/>{BORDERS}{shd}'
            f'<w:tcMar><w:top w:w="60" w:type="dxa"/><w:left w:w="100" w:type="dxa"/><w:bottom w:w="60" w:type="dxa"/><w:right w:w="100" w:type="dxa"/></w:tcMar>'
            f'</w:tcPr><w:p>{run}</w:p></w:tc>')


class Sheet:
    def __init__(self, base):
        self.d = docx.Document(base)
        body = self.d.element.body
        for el in list(body):
            if not el.tag.endswith("sectPr"):
                body.remove(el)
        xml = self.d.element.xml
        for rid, rel in list(self.d.part.rels.items()):
            if "image" in rel.reltype and f'"{rid}"' not in xml:
                self.d.part.drop_rel(rid)
        # header/footer come from HW1 ("HW1 – page N")
        for sec in self.d.sections:
            for part in (sec.header, sec.footer, sec.first_page_header, sec.first_page_footer):
                for t in part._element.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"):
                    if t.text and "HW1" in t.text:
                        t.text = t.text.replace("HW1", "HW2")

    def p(self, text="", bold=False, size=None, style=None, mono=False):
        para = self.d.add_paragraph(style=style)
        if text:
            r = para.add_run(text)
            r.bold = bold or None
            if size:
                r.font.size = Pt(size)
            if mono:
                r.font.name = "Consolas"
                r.font.size = Pt(10)
        return para

    def h(self, text, level):
        self.d.add_heading(text, level)

    def label_value(self, label, value="________"):
        para = self.d.add_paragraph()
        para.add_run(label).bold = True
        para.add_run(" " + value)

    def option(self, letter, text):
        para = self.d.add_paragraph()
        para.add_run(f"{letter}.  ")
        r = para.add_run(text)
        r.font.name = "Consolas"
        r.font.size = Pt(10)

    def box(self, height_in=1.6, hint=""):
        twips = int(height_in * 1440)
        run = (f'<w:r><w:rPr><w:i/><w:color w:val="808080"/><w:sz w:val="18"/></w:rPr>'
               f'<w:t xml:space="preserve">{escape(hint)}</w:t></w:r>') if hint else ""
        xml = (f'<w:tbl {W}>{TBL_PR}<w:tblGrid><w:gridCol w:w="9360"/></w:tblGrid>'
               f'<w:tr><w:trPr><w:cantSplit/><w:trHeight w:val="{twips}"/></w:trPr>'
               f'<w:tc><w:tcPr><w:tcW w:w="9360" w:type="dxa"/>{BORDERS}'
               f'<w:tcMar><w:top w:w="80" w:type="dxa"/><w:left w:w="120" w:type="dxa"/><w:bottom w:w="80" w:type="dxa"/><w:right w:w="120" w:type="dxa"/></w:tcMar>'
               f'</w:tcPr><w:p>{run}</w:p></w:tc></w:tr></w:tbl>')
        self._append(parse_xml(xml))
        self.p()

    def table(self, header, rows, widths=None, size=None):
        n = len(header)
        widths = widths or [9360 // n] * n
        grid = "".join(f'<w:gridCol w:w="{w}"/>' for w in widths)
        trs = '<w:tr><w:trPr><w:cantSplit/><w:tblHeader/></w:trPr>' + "".join(
            _cell(w, t, header=True, size=size) for w, t in zip(widths, header)) + "</w:tr>"
        for row in rows:
            row = list(row) + [""] * (n - len(row))
            trs += '<w:tr><w:trPr><w:cantSplit/></w:trPr>' + "".join(
                _cell(w, t, size=size) for w, t in zip(widths, row)) + "</w:tr>"
        self._append(parse_xml(f'<w:tbl {W}>{TBL_PR}<w:tblGrid>{grid}</w:tblGrid>{trs}</w:tbl>'))
        self.p()

    def _append(self, el):
        body = self.d.element.body
        body.insert(len(body) - 1, el)  # before sectPr

    def save(self, out):
        cp = self.d.core_properties
        cp.title = "Homework #2 (HW2) – Autumn 2026 GPU Computing"
        cp.subject = "HW2 answer sheet"
        self.d.save(out)


def build(base, out):
    s = Sheet(base)
    s.p("Homework #2 (HW2)", bold=True, size=18)
    s.p("Parallel Performance Theory, Gaussian Blur Kernels and More", size=13)
    s.p("Autumn 2026 – GPU Computing", size=11)
    s.table(["Field", ""], [["Name"], ["UW NetID"], ["Build environment"]], widths=[2400, 6960])

    # ---------------- Part 1 ----------------
    s.h("Part 1 – Questions", 1)
    mcq = [
        ("Assume that a kernel is launched with 1000 thread blocks each of which has 512 threads. If a variable is "
         "declared as a shared memory variable, how many versions of the variable will be created through the "
         "lifetime of the execution of the kernel?", ["1", "1,000", "512", "512,000"]),
        ("For a tiled matrix-matrix multiplication kernel, if we use a 32X32 tile, what is the reduction of memory "
         "bandwidth usage for input matrices A and B, compared to the non-tiled original?",
         ["1/8 of the original usage", "1/16 of the original usage", "1/32 of the original usage", "1/64 of the original usage"]),
        ("For the tiled single-precision matrix multiplication kernel as shown in lecture, assume that the tile size "
         "is 32X32 and the system has a DRAM burst size of 128 bytes. How many DRAM bursts will be delivered to the "
         "processor as a result of loading one A-matrix tile by a thread block?", ["16", "32", "64", "128"]),
    ]
    for i, (q, opts) in enumerate(mcq, 1):
        s.h(f"Question {i}", 2)
        s.p(q)
        for letter, o in zip("abcd", opts):
            s.option(letter, o)
        s.label_value("Answer:")
        s.p("Work / Reasoning:", bold=True)
        s.box(1.4)

    # ---------------- Part 2 ----------------
    s.h("Part 2 – Performance Theory Questions", 1)
    s.h("P1. Arithmetic Intensity of ML Operations", 2)
    s.p("AI := FLOPs / BYTE. FP32 unless stated; only global memory (DRAM) reads and writes count as Mem bytes; "
        "registers and cache are free. Transcendentals (exp, tanh, erf, sqrt) = 4 FLOP; FMA = 2 FLOP.")
    ops = [  # (title, hint, number of term rows)
        ("a. Matrix transpose (N×N matrix)", "B[j][i] = A[i][j]", 2),
        ("b. SoftMax, 3-pass, unfused", "pass 1: max   pass 2: exp + sum   pass 3: normalize", 3),
        ("c. Convolution (image W×H, kernel size K, channels C, batch size N)", "Out[n,c,h,w] = Σ In[...]·Wt[...]", 3),
    ]
    for title, hint, nterms in ops:
        s.h(title, 3)
        s.p("Algorithm equation and term definitions:", bold=True)
        s.box(1.2, hint)
        s.table(["algorithm", "FLOPs", "Mem bytes", "AI"], [[""]] * nterms + [["Total"]],
                widths=[3000, 2200, 2200, 1960])
        s.label_value("Order of complexity:", "________________")
        s.p()

    s.h("P2. Roofline Performance Charts", 2)
    s.table(["GPU", "Peak FP32 (TFLOP/s)", "Peak GMEM BW (GB/s)", "Ridge point (FLOP/B)"],
            [["Quadro RTX 6000 (TU102)"], ["H100 SXM5"]], widths=[3060, 2100, 2100, 2100])
    s.p("Chosen values for the variable terms in each AI expression:", bold=True)
    s.box(1.0)
    for g in ["a. Turing Quadro RTX 6000 (TU102)", "b. Hopper H100 SXM5"]:
        s.h(g, 3)
        s.p("Roofline plot:", bold=True)
        s.box(3.2)

    # ---------------- Part 3 ----------------
    s.h("Part 3 – Programming Exercises", 1)
    s.h("3.1 CUDA Gaussian Blur Kernels", 2)
    s.p("Blurred output images (test images #1, #2, #3):", bold=True)
    s.box(3.0)
    s.h("Extra Credit: Theoretical AI – global memory kernel vs shared memory kernel", 3)
    s.table(["Kernel", "FLOPs per output pixel", "GMEM bytes per output pixel", "AI (FLOP/B)"],
            [["Global memory"], ["Static shared memory"]], widths=[2400, 2300, 2400, 2260])
    s.p("Work / Reasoning:", bold=True)
    s.box(1.4)

    s.h("3.2 Host-Side Code Execution Timing", 2)
    s.p("Timer #1 = device alloc + HtoD copy + kernel + DtoH copy + cudaFree; Timer #2 = kernel launch → cudaDeviceSynchronize(). "
        "Not run in the debugger.")
    rows = []
    for img, size in [(1, "240×240"), (2, "512×512"), (3, "850×850"), (4, "4032×3024")]:
        for k in ["global", "static SMEM", "dynamic SMEM"]:
            rows.append([f"#{img} {size}", k])
    s.table(["Image", "Kernel", "Debug T1 (ms)", "Debug T2 (ms)", "Release T1 (ms)", "Release T2 (ms)"], rows,
            widths=[1700, 1660, 1500, 1500, 1500, 1500], size=18)
    s.p("DEBUG vs RELEASE comparison:", bold=True)
    s.box(1.2)

    s.h("3.3 Nsight VSE CUDA Debugger", 2)
    s.p("QUESTION: How many warps are running? Based on the Tesla T4 hardware, can you explain this number in terms of "
        "hardware resources?", bold=True)
    s.box(1.6)
    s.p("Screen capture 1 – Lanes + Warp Info, hw2_testimage1.png (240×240), full grid:", bold=True)
    s.box(3.2)
    s.p("Screen capture 2 – Lanes + Warp Info, <<<1,1>>> execution configuration:", bold=True)
    s.box(3.2)

    s.h("3.4 Nsight Compute CLI Profiling", 2)
    s.p("Kernel execution time = smsp__cycles_elapsed.sum / smsp__cycles_elapsed.sum.per_second (ns → s). "
        "Kernel instruction intensity = smsp__inst_executed.sum / dram__bytes.sum (Kbyte → byte).")
    rows = [[k, f"#{i}"] for k in ["global", "static SMEM"] for i in range(1, 5)]
    s.table(["Kernel", "Image", "smsp__cycles_ elapsed.sum", "…sum.per_ second", "smsp__inst_ executed.sum",
             "dram__bytes .sum", "Kernel time (s)", "Inst intensity (inst/B)"], rows,
            widths=[1100, 700, 1260, 1150, 1260, 1150, 1300, 1440], size=16)
    s.p("NCU CLI output (all 8 runs):", bold=True)
    s.box(3.0)
    s.h("HW Q: Why does kernel execution time differ for the shared mem version versus the global mem version?", 3)
    s.box(1.8)
    s.h("HW Q: Reason about observed differences between naive vs tiled (global vs shared) based on measured kernel "
        "execution times, instruction intensities and related memory and compute operations.", 3)
    s.box(1.8)

    s.h("Extra Credit: Host CPU Sequential Blur vs GPU Kernels", 2)
    s.table(["Image", "CPU T3 Debug (ms)", "CPU T3 Release (ms)", "GPU global T2 Release (ms)",
             "GPU static T2 Release (ms)", "GPU dynamic T2 Release (ms)"],
            [["#1 240×240"], ["#2 512×512"], ["#3 850×850"], ["#4 4032×3024"]],
            widths=[1560, 1560, 1560, 1560, 1560, 1560], size=18)
    s.save(out)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
